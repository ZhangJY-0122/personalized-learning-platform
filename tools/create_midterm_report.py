from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "template-reference.docx"
OUTPUT = ROOT / "1_专业综合课程设计_报告.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial Unicode MS"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, font_size=10.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_tr_pr.append(header_repeat)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="000000", align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size)
        set_cell_shading(table.rows[0].cells[i], "EDEDED")
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=font_size)
            if ri % 2 == 1:
                set_cell_shading(cells[i], "F8F8F8")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def clear_part(part):
    element = part._element
    for child in list(element):
        element.remove(child)
    part.add_paragraph()


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(24)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Arial Unicode MS"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.line_spacing = 1.2
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Arial Unicode MS"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.line_spacing = 1.2
    h2.paragraph_format.keep_with_next = True


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.name = "Arial Unicode MS"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_toc_line(doc, text):
    p = doc.add_paragraph(style="toc 1" if text[0].isdigit() else "toc 2")
    p.paragraph_format.left_indent = Inches(0.25 if text[0].isdigit() else 0.5)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Arial Unicode MS"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    r.font.size = Pt(10.5)
    return p


def build():
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    configure_styles(doc)

    # A4 is inherited from the reference template; normalize the working section.
    section = doc.sections[0]
    section.page_width = Inches(8.268)
    section.page_height = Inches(11.693)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.79)
    section.right_margin = Inches(0.79)
    section.different_first_page_header_footer = True
    clear_part(section.header)
    clear_part(section.first_page_header)
    clear_part(section.first_page_footer)
    clear_part(section.footer)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ""
    add_field(footer, "PAGE")

    # Cover page, matching the reference's cover structure.
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("【第1组】")
    r.font.name = "Arial Unicode MS"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    r.font.size = Pt(14)

    for _ in range(1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

    for text in ("计 算 机 专 业 综 合", "课 程 设 计 报 告"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = "Arial Unicode MS"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        r.font.size = Pt(25)
        r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("基于知识追踪与混合推荐算法的个性化学习平台")
    r.font.name = "Arial Unicode MS"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    r.font.size = Pt(14)
    r.bold = True

    cover = doc.add_table(rows=8, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = True
    set_table_borders(cover, color="BFBFBF", size="6")
    cover_data = [
        ("学院、系", "珠海科技学院计算机学院"),
        ("专业名称", "计算机科学与技术"),
        ("文档类型", "专业综合课程设计报告"),
        ("题目", "个性化学习平台"),
        ("组号", "第1组"),
        ("学生姓名", "张佳仪"),
        ("学号", "04230904"),
        ("班级", "PI230901"),
    ]
    for row, (a, b) in zip(cover.rows, cover_data):
        set_cell_text(row.cells[0], a, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
        set_cell_text(row.cells[1], b, align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
        set_cell_margins(row.cells[0], top=48, bottom=48, start=100, end=100)
        set_cell_margins(row.cells[1], top=48, bottom=48, start=100, end=100)
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(5.35)

    doc.add_page_break()

    # Contents page.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("目 录")
    r.font.name = "Arial Unicode MS"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    r.font.size = Pt(18)
    r.bold = True
    toc = [
        "题目",
        "1 引言",
        "2 需求分析",
        "3 系统设计",
        "4 数据库概念结构设计",
        "5 数据库逻辑结构设计",
        "6 安全性设计",
        "7 实现与阶段计划",
        "8 测试",
        "9 总结",
        "10 参考文献",
        "11 附录",
    ]
    for line in toc:
        add_toc_line(doc, line)

    doc.add_page_break()

    # Main report.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于知识追踪与混合推荐算法的个性化学习平台")
    r.font.name = "Arial Unicode MS"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    r.font.size = Pt(16)
    r.bold = True

    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 目的", 2)
    add_body(doc, "本报告总结个性化学习平台从完成选题、制定计划，到需求分析、总体设计、UI设计、数据设计、接口设计、实现和测试的过程，重点说明过程中遇到的问题、分析方法与解决方案。报告选择“契约先行的松耦合开发”作为重点剖析对象，说明如何在多人协作和功能持续扩展的情况下，保持前端、后端、数据库与推荐算法之间的边界清晰，并反思这种方法的优缺点。")
    add_heading(doc, "1.2 背景", 2)
    add_body(doc, "项目面向在线学习场景，目标是让学生能够登录课程、浏览资源、完成练习，并根据真实作答记录获得掌握度分析和个性化内容推荐。初期容易把页面展示、推荐结果和学习状态写成相互耦合的临时代码，后续一旦修改数据结构或算法，前端、接口和数据库会同时返工。因此本项目采用 Vue 3 前端、Spring Boot 服务、MySQL 数据库和增量迁移，先建立可运行的核心闭环，再逐步加入掌握度、推荐与反馈功能。")

    add_heading(doc, "2 需求分析", 1)
    add_heading(doc, "2.1 总体目标", 2)
    add_body(doc, "系统总体目标是形成“登录与授权—课程学习—题目作答—可信事件—掌握度更新—推荐反馈”的可追溯闭环。学生看到的掌握度和推荐结果必须来自后端保存的学习证据，而不是由前端写死或根据页面点击直接推断。期中阶段优先完成一门示例课程和电脑端核心流程，正式数据实验、教师统计和生产级部署作为后续目标。")
    add_heading(doc, "2.2 功能目标", 2)
    add_table(doc, ["用户类型", "主要功能", "权限边界"], [
        ("学生", "登录、课程目录、资源阅读、练习、掌握度、推荐与反馈", "只读取本人学习结果和已授权课程"),
        ("教师", "浏览授权课程和教学资源", "不能跨课程访问学生个人结果"),
        ("管理员", "课程维护、异常重试和数据恢复", "维护操作记录审计，按最小权限执行"),
    ], widths=[1.0, 3.4, 2.3])
    add_heading(doc, "2.3 数据需求", 2)
    add_body(doc, "核心数据包括用户与角色、课程与目录版本、章节与知识点、资源与题目、作答提交、学习事件、掌握度状态、历史记录、推荐批次和反馈。数据需求分析中最重要的判断是区分“行为数据”和“掌握证据”：浏览、点击和自报完成可以记录学习行为，但不能直接当成答对题目或掌握知识点的证据；只有服务端判分后的可信作答事件才能更新 BKT 状态。")
    add_heading(doc, "2.4 数据流图", 2)
    add_body(doc, "关键数据流可概括为：Vue 页面 → Nginx 同源代理 → Spring Boot 认证、授权与业务服务 → MySQL；服务端完成判分、事件入队、掌握度更新和推荐生成后，再以统一响应返回页面。推荐模块只读取版本化的学习状态和课程快照，反馈模块通过独立接口写入行为台账，从而避免 UI 层直接修改学习状态。")
    add_heading(doc, "2.5 性能要求", 2)
    add_body(doc, "当前阶段定位为本地单实例演示，性能要求以可验证和稳定为主：单事件处理目标 P95 不高于 300 ms，作答后页面可见状态目标 P95 不高于 2 s，单次本地推荐生成目标不高于 2 s。由于尚未完成正式压测，报告不把本地烟测结果当作生产容量结论；后续需要补充并发、重启、备份恢复和大数据量重建测试。")

    add_heading(doc, "3 系统设计", 1)
    add_heading(doc, "3.1 总体设计", 2)
    add_body(doc, "系统采用分层和模块化设计：前端负责交互与状态展示，后端负责认证、授权、业务写入和事务，数据库负责持久化，算法模块通过明确输入输出参与掌握度和推荐计算。项目计划按“开发准备—基础平台—学习状态—推荐反馈”的顺序推进，每一阶段设置接口、数据库和浏览器验收条件，只有上一阶段稳定后才引入下一组风险。")
    add_body(doc, "开发中最突出的问题是多人同时修改字段和接口，容易出现前端字段名、后端 DTO、数据库列名不一致。解决办法是先冻结 OpenAPI、统一 ID、错误码、分页和版本字段，所有跨模块写操作要求请求哈希和 Idempotency-Key；服务端拥有业务写入权，算法只接收结构化输入并返回结果，不直接操作业务表。这样实现了“可替换算法、可独立测试、可回放数据”的松耦合结构。")
    add_heading(doc, "3.2 UI设计", 2)
    add_body(doc, "UI设计围绕学生任务流组织为登录、课程目录、资源/练习、掌握度、推荐和历史等页面。设计时没有只画正常流程，而是同时定义加载、空数据、无权限、接口错误、处理中、过期和降级状态。实际遇到的问题是早期页面容易用模拟掌握度填充，视觉上完成了功能但无法证明数据链路有效；后续改为所有页面调用真实接口，页面只展示服务端返回的状态，并在推荐结果中显示原因、状态版本和失效提示。")
    add_heading(doc, "3.3 接口设计", 2)
    add_table(doc, ["接口", "用途", "松耦合约束"], [
        ("POST /auth/login", "登录并签发令牌", "统一错误码和 traceId，不暴露密码信息"),
        ("GET /courses/{courseId}/structure", "读取课程目录", "按课程授权和目录版本返回快照"),
        ("POST /practice/submissions", "提交练习", "服务端判分；幂等键绑定用户和请求哈希"),
        ("GET /students/{userId}/mastery", "读取掌握度", "只读查询，不在 GET 中隐式写入状态"),
        ("POST /students/{userId}/recommendations/generate", "生成推荐批次", "绑定 stateRevision、目录版本和策略版本"),
        ("POST /recommendations/{recommendationId}/feedback", "记录反馈", "行为台账与 BKT 状态分离"),
    ], widths=[2.7, 2.0, 2.0], font_size=9.5)

    add_heading(doc, "4 数据库概念结构设计", 1)
    add_body(doc, "系统概念模型包含五条主要关系：用户通过选课关系访问课程；课程由目录、章节和知识点组成，并关联资源和题目；学生提交作答后形成学习事件；学习事件按顺序更新掌握度、历史和画像；推荐批次关联推荐项目和反馈。目录、题目和推荐均保存版本信息，保证旧结果可以解释和复算。")
    add_table(doc, ["模块", "核心对象", "关系说明"], [
        ("课程模块", "course、catalog_version、chapter、knowledge_point", "课程拥有多个目录版本，章节和知识点属于指定版本"),
        ("学习模块", "submission、learning_event、mastery_state、mastery_history", "一次有效作答形成可信事件并产生状态历史"),
        ("推荐模块", "recommendation_batch、recommendation_item、feedback", "推荐批次绑定状态版本，反馈不直接更新掌握度"),
    ], widths=[1.2, 2.8, 2.7])

    add_heading(doc, "5 数据库逻辑结构设计", 1)
    add_heading(doc, "5.1 逻辑结构设计", 2)
    add_body(doc, "逻辑结构按关系模型落地，核心表使用 UUID 或稳定业务 ID，外键保证课程、用户、事件和推荐之间的引用一致。学习状态表保存当前值，历史表保存更新前后值、事件 ID 和模型版本；这样既能快速查询，也能在异常时按事件顺序重建。")
    add_heading(doc, "5.2 物理结构设计", 2)
    add_body(doc, "数据库采用 MySQL 8，使用 Flyway 进行增量迁移，已应用迁移不直接改写。针对用户课程、目录版本、事件顺序和推荐批次建立组合索引；对幂等键设置唯一约束。物理设计中的难点是既要保证查询速度，又要保留足够的审计信息，最终采用“当前状态 + 不可变历史 + 可重建事件”的折中方案。")

    add_heading(doc, "6 安全性设计", 1)
    add_heading(doc, "6.1 身份验证模式", 2)
    add_body(doc, "系统采用本地账号、BCrypt 密码哈希和 HS256 JWT。令牌校验签名、过期时间、签发者、受众和账号状态；退出时写入撤销记录。密码不以明文存储，题目接口也不返回标准答案。")
    add_heading(doc, "6.2 登录管理", 2)
    add_body(doc, "登录失败统一返回业务错误，不向客户端暴露账号是否存在；页面保存最小会话信息，服务端以令牌身份为准。后续还需要补充登录限流、HTTPS、密钥管理和审计日志，当前实现不宣称达到生产安全等级。")
    add_heading(doc, "6.3 权限管理", 2)
    add_body(doc, "授权在后端按用户、课程和资源对象逐级检查。学生只能读取本人数据，教师只能访问授权课程，管理员维护操作也需要明确角色。资源正文按纯文本渲染，减少将不可信内容当作 HTML 执行的风险。")

    add_heading(doc, "7 实现与阶段计划", 1)
    add_table(doc, ["阶段", "主要任务", "遇到的问题与解决"], [
        ("选题与计划", "确定个性化学习平台和分阶段目标", "范围过大；采用 M0 核心闭环，暂缓正式实验和复杂模型"),
        ("需求与总体设计", "冻结角色、数据边界、接口契约", "多人理解不一致；用 OpenAPI、统一 ID 和错误码约束协作"),
        ("UI与接口", "完成真实登录、课程、练习、掌握度和推荐页面", "模拟数据掩盖问题；改为真实 API，并覆盖空、错、处理中状态"),
        ("数据与实现", "完成事件、BKT、推荐和反馈持久化", "行为和掌握证据混淆；拆分行为台账与可信事件"),
        ("报告撰写", "整理问题、证据和后续边界", "容易把计划写成成果；明确区分已完成、演示边界和后续工作"),
    ], widths=[1.15, 2.7, 2.85], font_size=9.6)
    add_body(doc, "实现过程中采用 Vue 3、Vite、Pinia、Axios、Spring Boot 3、Java 17、MySQL 8 和 Flyway。BKT 与规则基线由 Java 在线计算，离线模型训练单独保留，避免算法服务故障阻塞基础学习流程。推荐模块先实现可解释排序和反馈闭环，再考虑模型替换。")

    add_heading(doc, "8 测试", 1)
    add_heading(doc, "8.1 测试技术", 2)
    add_body(doc, "测试采用接口测试、数据库断言、前端浏览器验收和故障场景复现相结合的方式。接口测试关注权限、字段和错误码；数据库测试关注迁移、幂等、事务回滚和状态版本；浏览器测试关注真实页面、刷新恢复和控制台错误。")
    add_heading(doc, "8.2 关键测试结果", 2)
    add_table(doc, ["场景", "验证结论"], [
        ("重复提交与网络重试", "同一 Idempotency-Key 只产生一条提交和一条学习事件；同键不同内容返回冲突"),
        ("事务失败与重试", "判分结果保留，学习状态整体回滚，失败事件可以按退避规则重试"),
        ("越权和答案保护", "未授权课程返回 403，题目接口不返回标准答案"),
        ("推荐失效", "学习状态或目录版本变化后旧批次标记失效，不伪装为最新结果"),
        ("浏览器闭环", "真实登录、作答、掌握度查询、推荐生成和反馈流程可完成，页面无未处理错误"),
    ], widths=[2.1, 4.6], font_size=9.8)

    add_heading(doc, "9 总结", 1)
    add_heading(doc, "9.1 心得体会", 2)
    add_body(doc, "本次课程设计中印象最深的是松耦合并不等于简单拆成几个文件，而是要明确每个模块拥有的数据、承担的责任和对外承诺。契约先行、服务端单一写入、事件可追溯和接口幂等，使前端 UI、后端事务和推荐算法可以分别演进；当需求变化或故障出现时，也能通过测试和历史数据定位影响范围。它的优点是协作并行度高、模块可替换、测试边界清晰，缺点是前期需要投入较多时间维护契约、版本和适配代码，跨模块问题的排查也更依赖日志和证据链。对于期中阶段而言，这种方法降低了后续返工风险，但不能代替正式性能、安全和教育效果评估。")
    add_heading(doc, "9.2 建议", 2)
    add_body(doc, "后续应继续完善三方面工作：第一，增加自动化契约测试，确保 OpenAPI 与真实响应同步；第二，扩充真实课程和公开数据，区分演示数据与研究数据；第三，补充并发压测、备份恢复、HTTPS、登录限流和安全审计。只有在数据质量、接口稳定性和非功能指标都得到验证后，才适合评价推荐算法的实际教学价值。")

    add_heading(doc, "10 参考文献", 1)
    add_body(doc, "[1] 李军国，吴昊，郭晓燕，王舒. 软件工程案例教程（第2版）[M]. 北京：清华大学出版社，2013.")
    add_body(doc, "[2] Spring Boot Reference Documentation[EB/OL]. https://docs.spring.io/spring-boot/docs/current/reference/html/")
    add_body(doc, "[3] Vue.js Guide[EB/OL]. https://cn.vuejs.org/guide/")
    add_body(doc, "[4] Corbett A T, Anderson J R. Knowledge tracing: Modeling the acquisition of procedural knowledge[J]. User Modeling and User-Adapted Interaction, 1995.")

    add_heading(doc, "11 附录", 1)
    add_body(doc, "本报告对应的 OpenAPI 契约、数据库迁移、接口验收和浏览器测试证据均保存在课程设计工作区；标注为“后续”的内容不计入本阶段成果。")

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
